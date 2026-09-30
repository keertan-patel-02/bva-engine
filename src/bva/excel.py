
import pandas as pd
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment
from bva.variance import load_pnl, build_variance, flag_material, summarize, load_drivers
from bva.bridge import build_bridge


OUTPUT_DIR = Path(__file__).resolve().parents[2] / "output"
MONEY = "#,##0.00;(#,##0.00)"
PERCENT = "0.0%"
LABELS = {
    "section": "Section",
    "line_item": "Line Item",
    "budget": "Budget",
    "actual": "Actual",
    "variance": "Variance",
    "variance_pct": "Variance %",
    "variance_fav": "Variance Favorable",
    "material": "Material",
    "component": "Component",
    "amount": "Amount",
}
FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")

def output_excel(period):
    """
    Creates an excel file with a summary table for the property portfolio and variance tables and bridge tables per property for a given period.
    Arguments -> period
    Requires that period be in the PNL.
    """
    pnl = load_pnl()
    available = sorted(pnl["period"].unique())
    if not (pnl["period"] == period).any():
        raise ValueError(f"{period!r} not in the P&L. Available periods: {', '.join(available)}")
    drivers = load_drivers()
    summary_df = summarize(build_variance(pnl, period))
    path=OUTPUT_DIR/f"bva_pack_{period}.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary_df = summary_df.rename(columns=LABELS)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        
        format_sheet(writer.sheets["Summary"])

        for entity in sorted(pnl["entity_id"].unique()):
            entity_df = flag_material(build_variance(pnl, period, entity))
            display_df = entity_df[["section", "line_item", "budget", "actual", "variance", "variance_pct", "variance_fav"]]
            display_df = display_df.rename(columns=LABELS)
            entity_bridge_df = build_bridge(drivers, period, entity)
            entity_bridge_df = entity_bridge_df[["component", "amount"]]
            entity_bridge_df = entity_bridge_df.rename(columns=LABELS)
            display_df.to_excel(writer, sheet_name=entity, index=False)
            format_sheet(writer.sheets[entity])
            highlight_material(writer.sheets[entity], entity_df)
            entity_bridge_df.to_excel(writer, sheet_name=f"{entity}_bridge", index=False)
            format_sheet(writer.sheets[f"{entity}_bridge"])
    
    return path
        

def format_sheet(ws):
    ws.freeze_panes ="A2"
    ## row 1 holds the headers
    for header_cell in ws[1]:
        name = header_cell.value
        letter = header_cell.column_letter

        header_cell.font = Font(bold=True)

        if name in ("Budget", "Actual", "Variance", "Variance Favorable", "Amount"):
            fmt = MONEY
            ws.column_dimensions[letter].width = 18
        elif name == "Variance %":
            fmt = PERCENT
            ws.column_dimensions[letter].width = 10
        else:
            ws.column_dimensions[letter].width = max(len(str(cell.value)) for cell in ws[letter]) + 2
            continue

        for cell in ws[letter][1:]:
            cell.number_format = fmt

def highlight_material(ws, df):
    flag = int(df["material"].sum())

    for row_number, is_material in enumerate(df["material"], start=2):
        if is_material:
            for cell in ws[row_number]:
                cell.fill = FILL
    if flag:
        note = f"Highlighted rows exceed both materiality thresholds: $2,500 and 5%. Total Line(s): {flag}"
    else:
        note = "No lines exceeded both materiality thresholds ($2,500 and 5%)."

    ws.cell(row=ws.max_row + 2, column=1).value = note